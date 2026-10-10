"""Pitch deck authored from actual measured prototype results and screenshots."""
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
import json
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'submission'/'TrustBoundary_Hackathon_Pitch.pptx'
M=json.loads((ROOT/'reports'/'metrics.json').read_text())
r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
BG='081322';PANEL='112238';PANEL2='162c43';TEAL='39D9B0';WHITE='EEF7FF';MUTED='9BB0C8';RED='FF8189';GOLD='F1C174';BLUE='85A9FF'

def rgb(s):return RGBColor.from_string(s)
def rect(sl,x,y,w,h,color,radius=False,edge=None):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(color)
 sh.line.fill.background() if not edge else None
 if edge:sh.line.color.rgb=rgb(edge)
 return sh

def text(sl,msg,x,y,w,h,size=16,color=WHITE,bold=False,spacing=1,align=None,font='Aptos',valign=None):
 box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=Inches(.04);tf.margin_top=Inches(.01);tf.margin_right=Inches(.04);tf.margin_bottom=Inches(.01)
 p=tf.paragraphs[0];p.text=msg;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(spacing)
 if align:p.alignment=align
 if valign:tf.vertical_anchor=valign
 return box

def bg(sl,n,kicker,title,subtitle=''):
 sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(BG)
 rect(sl,.35,.42,.08,.3,TEAL)
 text(sl,'TRUSTBOUNDARY / '+kicker.upper(),.54,.40,7,.32,10,TEAL,True)
 text(sl,f'{n:02d} / 10',11.7,.43,1.25,.32,10,MUTED,True,align=PP_ALIGN.RIGHT)
 text(sl,title,.57,1.00,12.1,.7,29,WHITE,True)
 if subtitle:text(sl,subtitle,.59,1.78,11.9,.58,13,MUTED)
 rect(sl,.55,7.16,12.2,.008,'304259')
 text(sl,'ET × ACCENTURE AI HACKATHON 2026 · AGENTIC EDITION',.61,7.2,8,.2,8,MUTED)
 text(sl,'v1.2 · F3–D1 · SYNTHETIC SANDBOX',9.2,7.2,3.5,.2,8,MUTED,align=PP_ALIGN.RIGHT)

def new(n,k,t,s=''):
 sl=r.slides.add_slide(r.slide_layouts[6]);bg(sl,n,k,t,s);return sl

def card(sl,x,y,w,h,eyebrow,value,desc,color=TEAL):
 rect(sl,x,y,w,h,PANEL,True)
 text(sl,eyebrow,x+.2,y+.2,w-.35,.27,10,MUTED,True)
 text(sl,value,x+.19,y+.55,w-.35,.58,26,color,True)
 text(sl,desc,x+.19,y+1.20,w-.35,h-1.30,11,WHITE)

def screenshot(sl,path,x,y,w,h):
 from PIL import Image
 im=Image.open(path);iw,ih=im.size;ratio=iw/ih;boxratio=w/h
 # crop to slide rectangle: favor top and center, show real UI, no graphical edits.
 if ratio>boxratio:
  crop=(1-boxratio/ratio)/2
  pic=sl.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h));pic.crop_left=crop;pic.crop_right=crop
 else:
  crop=(1-ratio/boxratio)
  pic=sl.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h));pic.crop_top=0;pic.crop_bottom=crop

# 1 cover
sl=new(1,'mission','Enterprise agents can act. Can they be trusted?','Prompt Injection Firewall · Problem Statement #2')
rect(sl,.7,2.5,11.9,3.5,PANEL,True)
text(sl,'TRUSTBOUNDARY',1.05,2.8,10,.84,46,WHITE,True)
text(sl,'AI',10.75,2.8,1.0,.83,46,TEAL,True)
text(sl,'THE AGENTIC SECURITY CONTROL PLANE',1.08,3.95,10,.46,19,TEAL,True)
text(sl,'Detect hostile instructions · Contain unauthorized tool actions · Measure true task utility',1.08,4.65,10.6,.73,17,MUTED)

# 2 Problem
sl=new(2,'why now','Prompt injection becomes business risk when AI can act.','An email, retrieved page or tool response may contain a forged high-priority instruction.')
for x,h,t,d in [(0.72,'01 / ATTACK','Untrusted input','Merchant email hides a role override or encoded instruction.'),(4.88,'02 / ESCALATION','Tool manipulation','Agent is pushed toward refund, secret or outbound-message tool.'),(9.04,'03 / CONSEQUENCE','Business harm','Unauthorized financial action, leaked secrets, or interrupted workflow.')]:
 rect(sl,x,2.65,3.55,2.9,PANEL,True)
 text(sl,h,x+.23,2.96,3.0,.3,11,TEAL,True)
 text(sl,t,x+.23,3.46,3.1,.6,21,WHITE,True)
 text(sl,d,x+.23,4.24,3.08,.95,15,MUTED)
text(sl,'Product principle: External content supplies data, never authority.',.82,6.25,11.7,.45,20,TEAL,True)

# 3 solution
sl=new(3,'product','Not a chatbot. A security layer around agent actions.','Three bounded components plus independent authorization and EGO measurement.')
for y,title,descr in [(2.45,'INSPECT','Normalize sources, inspect risky spans, assign provenance & attack categories.'),(3.55,'ENFORCE','Deterministic role + source-trust gate; privileged tools need separate approval.'),(4.65,'EVALUATE','Audit tool outcomes, triage incidents and replay reproducible security tests.')]:
 rect(sl,.86,y,7.25,.86,PANEL,True)
 text(sl,title,1.08,y+.18,1.8,.38,17,TEAL,True)
 text(sl,descr,2.92,y+.16,4.95,.48,13,WHITE)
card(sl,8.46,2.45,3.68,3.7,'DEFENSE-IN-DEPTH','9 categories','Optional hosted LLM + locally trained TF-IDF risk model + rule layer; authorization stays outside the model.',BLUE)

# 4 screenshot lab
sl=new(4,'working prototype','Live attack laboratory and case comparison.','Same user task, same untrusted message, different tool authorization outcomes.')
screenshot(sl,ROOT/'submission'/'screenshots'/'02_before_after.png',.62,2.44,9.25,4.23)
rect(sl,10.05,2.44,2.62,4.23,PANEL,True)
text(sl,'ROLE SPOOFING',10.22,2.76,2.2,.35,13,TEAL,True)
text(sl,'Insecure scripted control follows the injected refund request.',10.22,3.32,2.24,1.04,14,RED)
text(sl,'Protected agent completes the ticket and denies the external-origin refund.',10.22,4.6,2.24,1.53,14,WHITE)

# 5 architecture
sl=new(5,'architecture','Provenance → inspection → policy → tool authorization.','The classifier is advisory; the tool gateway is the actual safety boundary.')
items=[('USER TASK','trusted goal'),('CONTENT INTAKE','email · HTML · PDF · API'),('INSPECTION COMPONENT','local ML + rules + optional LLM'),('RISK DISPOSITION','allow / sanitize / quarantine / escalate'),('BUSINESS AGENT','settlement read · ticket draft'),('DETERMINISTIC TOOL GATE','role · source · independent approval'),('AUDIT + EGO','trace · score · incident triage')]
for i,(a,b) in enumerate(items):
 y=2.24+i*.59;rect(sl,2.4,y,8.6,.50,PANEL2 if i in [3,5] else PANEL,True,TEAL if i in [3,5] else None)
 text(sl,a,2.6,y+.08,3.6,.26,12,TEAL if i in [3,5] else WHITE,True)
 text(sl,b,6.0,y+.09,4.7,.24,10,MUTED)
 if i<6:text(sl,'↓',6.35,y+.43,.7,.24,13,TEAL,True,align=PP_ALIGN.CENTER)

# 6 category coverage
sl=new(6,'official challenge','Nine attack patterns implemented and unit-tested.','Official F3 requires ≥7 tested categories. Broad real-world coverage remains unverified.')
names=['Instruction Override','Role Change','Secret Extraction','Tool Abuse','Credential Theft','Context Poisoning','Multi-Step Jailbreaks','Encoded Instructions','Indirect Prompt Injection']
for idx,name in enumerate(names):
 col=idx%3;row=idx//3;x=.8+col*4.2;y=2.42+row*1.15
 rect(sl,x,y,3.82,.92,PANEL,True)
 text(sl,f'{idx+1:02d}',x+.16,y+.25,.4,.35,14,TEAL,True)
 text(sl,name,x+.70,y+.24,2.87,.5,14,WHITE,True)
text(sl,'Evidence: labeled fixtures · regression tests · per-category synthetic holdout · redacted spans',.87,6.26,11.9,.35,15,MUTED)

# 7 measured benchmark
sl=new(7,'measured evidence','Synthetic holdout: reproducible, not production validation.','Local hybrid detector measured against 80 synthetic held-out examples, with transparent caveats.')
nums=[('ATTACKS',str(M['attack_cases']),'test examples'),('BENIGN / QUOTED',str(M['benign_and_ambiguous_cases']),'test examples'),('DETECTION RECALL',f"{M['recall']:.0%}",'templated examples'),('FALSE POSITIVES',str(M['fp']),'held-out cases')]
for i,(k,v,d) in enumerate(nums):card(sl,.75+i*3.14,2.37,2.86,1.99,k,v,d,TEAL if i in [0,2] else BLUE)
rect(sl,.75,4.75,11.85,1.52,PANEL2,True)
text(sl,f"Forbidden protected executions: {M['unauthorized_protected_tool_executions']}   •   Scripted control ASR: {M['scripted_baseline_attack_success_rate']:.1%}   •   Protected ASR: {M['protected_attack_success_rate']:.1%}",.97,5.06,11.3,.42,17,WHITE,True)
text(sl,'Important: The control is intentionally vulnerable scripted logic, NOT a measured real LLM. Dataset is templated; a perfect score does not prove security.',.99,5.65,11.12,.47,11,GOLD)

# 8 EGO screen
sl=new(8,'quality engineering','EGO measures safety AND usefulness.','Security that breaks every legitimate workflow is not a viable enterprise product.')
screenshot(sl,ROOT/'submission'/'screenshots'/'03_evaluation.png',.64,2.35,8.1,4.43)
for y,a,b in [(2.62,'EVALUATION','Attack recall, precision, per-class evidence'),(4.00,'GUARDRAILS','No unauthorized mock tool execution'),(5.37,'OBSERVABILITY','Audit chain, latency and failure analysis')]:
 rect(sl,9.03,y,3.55,1.07,PANEL,True)
 text(sl,a,9.24,y+.14,3.08,.3,14,TEAL,True)
 text(sl,b,9.24,y+.49,3.08,.48,12,MUTED)

# 9 leadership impact
sl=new(9,'product + business','Designed like an enterprise AI platform, not a toy.','A synthetic fintech operations scenario connects risk to real business decision-making.')
for x,h,b,d in [(.8,'CONTROL PLANE','Platform governance','Central AI policy and model-independent safety checks.'),(4.97,'AGENT STUDIO','Orchestration','Inspection, task execution, incident analysis and replay.'),(9.10,'PRODUCT KPIs','Measurable ROI','Safety, false positives, permitted-workflow success, costs.')]:
 rect(sl,x,2.49,3.48,2.97,PANEL,True)
 text(sl,h,x+.2,2.79,3.0,.27,11,TEAL,True)
 text(sl,b,x+.2,3.25,3.06,.58,22,WHITE,True)
 text(sl,d,x+.2,4.09,3.07,.91,15,MUTED)
text(sl,'Financial-impact dashboard uses editable hypothetical assumptions; no real customer savings claimed.',.87,6.12,11.6,.48,16,GOLD)

# 10 roadmap and judge ask
sl=new(10,'submission','Ready to demonstrate. Honest about the next steps.','Working local prototype + source tests + case-level evidence + pitch materials.')
for idx,(label,desc) in enumerate([('TODAY','Run the protected fintech workflow, evaluate, inspect traces.'),('NEXT','External adversarial corpus, LLM red-team runs, real identity/approval integration.'),('FOR PRODUCTION','Multitenant auth, audit hardening, SIEM, OCR, rate limits, independent security review.')]):
 y=2.42+idx*1.22;rect(sl,.89,y,11.45,.99,PANEL,True);text(sl,label,1.16,y+.25,2.6,.34,14,TEAL,True);text(sl,desc,3.57,y+.23,8.35,.54,14,WHITE)
text(sl,'TrustBoundary: provenance before permission.',1.1,6.37,10,.45,22,TEAL,True)
r.save(OUT)
print(OUT)
